"""Deterministic projection from canonical Slackpipe data."""

from __future__ import annotations

import csv
import hashlib
import io
import json
from collections import Counter
from contextlib import suppress
from pathlib import Path
from typing import Any

import duckdb
from docx import Document
from pptx import Presentation
from pypdf import PdfReader

from slackquery.db import connect_state, path_sql
from slackquery.models import ProjectionStats
from slackquery.settings import Settings

CANONICAL_SCHEMA_VERSION = 2
MESSAGE_RECIPE_VERSION = "message-v1"
THREAD_RECIPE_VERSION = "thread-context-v1"
FILE_RECIPE_VERSION = "file-chunk-v1"
TEXT_RECIPE_VERSION = MESSAGE_RECIPE_VERSION
_TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".rst",
    ".log",
    ".json",
    ".jsonl",
    ".xml",
    ".yaml",
    ".yml",
    ".toml",
    ".ini",
    ".cfg",
    ".conf",
    ".py",
    ".js",
    ".jsx",
    ".ts",
    ".tsx",
    ".java",
    ".go",
    ".rs",
    ".rb",
    ".php",
    ".sh",
    ".bash",
    ".zsh",
    ".sql",
    ".css",
    ".scss",
    ".html",
    ".htm",
    ".c",
    ".h",
    ".cpp",
    ".hpp",
    ".cs",
    ".swift",
    ".kt",
    ".kts",
    ".scala",
    ".lua",
    ".r",
}


def _watermark(connection: duckdb.DuckDBPyConnection) -> str:
    values = connection.execute(
        """
        SELECT count(*), coalesce(max(last_changed_at)::VARCHAR, ''),
               coalesce(max(latest_run_id), ''), coalesce(max(payload_hash), '')
        FROM canonical.messages
        """
    ).fetchone()
    assert values is not None
    file_value = ""
    if _table_exists(connection, "files"):
        columns = _columns(connection, "files")
        hash_column = next(
            (item for item in ("payload_hash", "content_hash", "sha256") if item in columns), None
        )
        path_column = "local_object_path" if "local_object_path" in columns else None
        expressions = ["count(*)::VARCHAR"]
        if hash_column:
            expressions.append(f"coalesce(max({hash_column}), '')")
        if path_column:
            expressions.append(f"coalesce(max({path_column}), '')")
        row = connection.execute(f"SELECT {', '.join(expressions)} FROM canonical.files").fetchone()
        file_value = "|".join(map(str, row or ()))
    return hashlib.sha256(("|".join(map(str, values)) + "|" + file_value).encode()).hexdigest()


def _table_exists(connection: duckdb.DuckDBPyConnection, name: str) -> bool:
    row = connection.execute(
        "SELECT count(*) FROM information_schema.tables "
        "WHERE table_catalog='canonical' AND table_name=?",
        [name],
    ).fetchone()
    assert row is not None
    return bool(row[0])


def _columns(connection: duckdb.DuckDBPyConnection, table: str) -> set[str]:
    return {
        row[0]
        for row in connection.execute(
            "SELECT column_name FROM information_schema.columns "
            "WHERE table_catalog='canonical' AND table_name=?",
            [table],
        ).fetchall()
    }


def _safe_attachment_path(root: Path, workspace_slug: str, local_path: str) -> Path:
    """Resolve an attachment while rejecting absolute paths and symlink traversal."""
    relative = Path(local_path)
    if relative.is_absolute() or ".." in relative.parts:
        raise ValueError("attachment path escapes workspace root")
    workspace_root = (root / workspace_slug).resolve()
    candidate = (workspace_root / relative).resolve(strict=True)
    try:
        candidate.relative_to(workspace_root)
    except ValueError as error:
        raise ValueError("attachment path escapes workspace root") from error
    if not candidate.is_file():
        raise ValueError("attachment is not a regular file")
    return candidate


def _attachment_candidates(
    settings: Settings,
    workspace_id: str,
    workspace_slug: str,
    workspace_name: str | None,
    local_path: str,
) -> list[Path]:
    """Resolve only explicit root/workspace combinations; never recursively search."""
    roots = [settings.attachment_root, *settings.attachment_search_roots]
    configured = settings.attachment_workspace_map.get(workspace_id)
    workspace_names = [configured, workspace_id, workspace_slug, workspace_name]
    result: list[Path] = []
    seen: set[Path] = set()
    for root in roots:
        try:
            resolved_root = root.resolve(strict=True)
        except FileNotFoundError:
            continue
        for workspace in workspace_names:
            if not workspace:
                continue
            try:
                candidate = _safe_attachment_path(resolved_root, workspace, local_path)
            except (FileNotFoundError, ValueError):
                continue
            if candidate not in seen:
                seen.add(candidate)
                result.append(candidate)
    return result


def extract_file_text(path: Path, mime_type: str | None, max_bytes: int) -> str:
    """Extract bounded text from one supported local attachment."""
    if path.stat().st_size > max_bytes:
        raise OverflowError("attachment exceeds configured byte limit")
    suffix = path.suffix.lower()
    mime = (mime_type or "").split(";", 1)[0].lower()
    if suffix == ".pdf" or mime == "application/pdf":
        return "\n\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages).strip()
    if suffix == ".docx" or mime.endswith("wordprocessingml.document"):
        document = Document(str(path))
        return "\n".join(paragraph.text for paragraph in document.paragraphs).strip()
    if suffix == ".pptx" or mime.endswith("presentationml.presentation"):
        presentation = Presentation(str(path))
        values: list[str] = []
        for slide in presentation.slides:
            values.extend(
                shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text
            )
        return "\n".join(values).strip()
    raw = path.read_bytes()
    if b"\x00" in raw[:8192]:
        raise ValueError("binary attachment rejected")
    if suffix == ".csv" or mime in {"text/csv", "application/csv"}:
        decoded = raw.decode("utf-8-sig")
        output = io.StringIO()
        writer = csv.writer(output, lineterminator="\n")
        writer.writerows(csv.reader(io.StringIO(decoded)))
        return output.getvalue().strip()
    if mime.startswith("text/") or suffix in _TEXT_EXTENSIONS:
        return raw.decode("utf-8-sig").strip()
    raise TypeError("unsupported attachment type")


def chunk_text(text: str, *, target_chars: int, overlap_chars: int) -> list[str]:
    """Produce stable character-approximated chunks with deterministic overlap."""
    if target_chars <= 0 or overlap_chars < 0 or overlap_chars >= target_chars:
        raise ValueError("invalid chunk bounds")
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not normalized:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(normalized):
        end = min(start + target_chars, len(normalized))
        if end < len(normalized):
            boundary = max(
                normalized.rfind("\n", start + target_chars // 2, end),
                normalized.rfind(" ", start + target_chars // 2, end),
            )
            if boundary > start:
                end = boundary
        value = normalized[start:end].strip()
        if value:
            chunks.append(value)
        if end >= len(normalized):
            break
        start = max(start + 1, end - overlap_chars)
    return chunks


def _insert_document(connection: duckdb.DuckDBPyConnection, values: list[Any]) -> None:
    columns = (
        "document_id, document_kind, workspace_id, workspace_name, workspace_slug, "
        "workspace_url, channel_id, channel_name, author_id, author_name, ts, ts_us, "
        "timestamp, thread_id, thread_root_id, is_thread_parent, is_thread_reply, "
        "is_thread_broadcast, text_display, text_lexical, text_embedding, source_version, "
        "text_hash, metadata_json, permalink, is_active, projected_at"
    )
    connection.execute(
        f"INSERT OR REPLACE INTO document_projection ({columns}) "
        f"VALUES ({', '.join('?' for _ in values)})",
        values,
    )


def _set_projection_identity(
    connection: duckdb.DuckDBPyConnection,
    document_id: str,
    source_identity: str,
    source_revision: str,
    embedding_recipe: str,
) -> None:
    connection.execute(
        """
        UPDATE document_projection
        SET source_identity=?, source_revision=?, embedding_recipe=?
        WHERE document_id=?
        """,
        [source_identity, source_revision, embedding_recipe, document_id],
    )


def _project_threads(
    connection: duckdb.DuckDBPyConnection, source_version: str, now: Any, settings: Settings
) -> int:
    rows = connection.execute(
        """
        SELECT document_id, workspace_id, workspace_name, workspace_slug, workspace_url,
          channel_id, channel_name, author_id, author_name, ts, ts_us, timestamp, thread_id,
          thread_root_id, text_display, permalink
        FROM document_projection
        WHERE source_identity=? AND document_kind='message' AND is_active
          AND thread_id IS NOT NULL
        ORDER BY thread_id, ts_us, document_id
        """,
        [source_version],
    ).fetchall()
    grouped: dict[str, list[tuple[Any, ...]]] = {}
    for row in rows:
        grouped.setdefault(row[12], []).append(row)
    count = 0
    for thread_id, messages in sorted(grouped.items()):
        blocks: list[str] = []
        used = 0
        for row in messages[: settings.thread_context_max_messages]:
            block = f"[{row[11].isoformat()}] {row[8] or row[7] or 'unknown'}: {row[14]}"
            addition = len(block) + (2 if blocks else 0)
            if used + addition > settings.thread_context_max_chars:
                break
            blocks.append(block)
            used += addition
        if not blocks:
            continue
        first = messages[0]
        text = "\n\n".join(blocks)
        digest = hashlib.sha256(text.encode()).hexdigest()
        document_id = f"thread:{thread_id}"
        metadata = {
            "thread_id": thread_id,
            "message_ids": [row[0] for row in messages[: len(blocks)]],
            "message_count": len(blocks),
        }
        _insert_document(
            connection,
            [
                document_id,
                "thread_context",
                first[1],
                first[2],
                first[3],
                first[4],
                first[5],
                first[6],
                first[7],
                first[8],
                first[9],
                first[10],
                first[11],
                thread_id,
                first[13],
                True,
                False,
                False,
                text,
                text,
                text,
                source_version,
                digest,
                json.dumps(metadata, sort_keys=True),
                first[15],
                True,
                now,
            ],
        )
        _set_projection_identity(
            connection, document_id, source_version, digest, THREAD_RECIPE_VERSION
        )
        count += 1
    return count


def _project_files(
    connection: duckdb.DuckDBPyConnection, source_version: str, now: Any, settings: Settings
) -> Counter[str]:
    stats: Counter[str] = Counter()
    if not _table_exists(connection, "files"):
        stats["files_table_missing"] += 1
        return stats
    columns = _columns(connection, "files")
    required = {"local_object_path", "workspace_id"}
    if not required.issubset(columns):
        stats["files_schema_unsupported"] += 1
        return stats

    def expression(candidates: tuple[str, ...], default: str = "NULL") -> str:
        found = next((name for name in candidates if name in columns), None)
        return f"f.{found}" if found else default

    file_id = expression(("file_id", "id", "file_key"), "f.local_object_path")
    message_key = expression(("message_key", "source_message_key"))
    message_ts = expression(("message_ts", "ts"))
    channel_id = expression(("channel_id",))
    mime = expression(("mimetype", "mime_type"))
    name = expression(("name", "filename", "title"), "f.local_object_path")
    deleted = expression(("is_deleted",), "false")
    rows = connection.execute(f"""
        SELECT f.workspace_id, coalesce(w.workspace_slug, f.workspace_id),
          coalesce(w.team_name, f.workspace_id), w.workspace_url,
          {file_id}, f.local_object_path, {name}, {mime}, {message_key}, {message_ts},
          {channel_id}, {deleted}, f.workspace_id
        FROM canonical.files f
        LEFT JOIN canonical.workspaces w ON w.workspace_id=f.workspace_id
        ORDER BY f.workspace_id, {file_id}, f.local_object_path
    """).fetchall()
    target_chars = settings.file_chunk_tokens * settings.token_char_approximation
    overlap_chars = settings.file_chunk_overlap_tokens * settings.token_char_approximation
    for row in rows:
        if row[11]:
            stats["deleted"] += 1
            continue
        if not row[5]:
            stats["local_path_missing"] += 1
            continue
        try:
            relative = Path(row[5])
            if relative.is_absolute() or ".." in relative.parts:
                raise ValueError("attachment path escapes workspace root")
            paths = _attachment_candidates(settings, row[12], row[1], row[2], row[5])
            if not paths:
                raise FileNotFoundError(row[5])
            path = paths[0]
            text = extract_file_text(path, row[7], settings.attachment_max_bytes)
            chunks = chunk_text(text, target_chars=target_chars, overlap_chars=overlap_chars)
        except FileNotFoundError:
            stats["missing"] += 1
            continue
        except OverflowError:
            stats["oversized"] += 1
            continue
        except TypeError:
            stats["unsupported"] += 1
            continue
        except (ValueError, OSError, UnicodeError):
            stats["rejected"] += 1
            continue
        message = None
        mapped_message_id = row[8]
        if row[8]:
            message = connection.execute(
                """
                SELECT channel_id, channel_name, author_id, author_name, ts, ts_us, timestamp,
                  thread_id, thread_root_id, permalink
                FROM document_projection WHERE document_id=? AND document_kind='message'
            """,
                [row[8]],
            ).fetchone()
        if message is None and row[9] and row[10]:
            mapped = connection.execute(
                """
                SELECT document_id, channel_id, channel_name, author_id, author_name, ts,
                  ts_us, timestamp, thread_id, thread_root_id, permalink
                FROM document_projection
                WHERE document_kind='message' AND is_active AND workspace_id=?
                  AND channel_id=? AND ts=?
                ORDER BY document_id LIMIT 1
                """,
                [row[12], row[10], row[9]],
            ).fetchone()
            if mapped is not None:
                mapped_message_id = mapped[0]
                message = mapped[1:]
        if message is None:
            channel = row[10] or "files"
            ts_us = 0
            timestamp = "1970-01-01T00:00:00+00:00"
            message = (channel, None, None, None, "0.000000", ts_us, timestamp, None, None, None)
        for index, chunk in enumerate(chunks):
            document_id = f"file:{row[0]}:{row[4]}:{index:06d}"
            digest = hashlib.sha256(chunk.encode()).hexdigest()
            metadata = {
                "file_id": str(row[4]),
                "file_name": row[6],
                "message_id": mapped_message_id,
                "chunk_index": index,
                "chunk_count": len(chunks),
                "mime_type": row[7],
            }
            embedding = (
                f"workspace: {row[1]}\nchannel: {message[1] or message[0]}\n"
                f"file: {row[6]}\ncontent: {chunk}"
            )
            _insert_document(
                connection,
                [
                    document_id,
                    "file_chunk",
                    row[0],
                    row[2],
                    row[1],
                    row[3],
                    message[0],
                    message[1],
                    message[2],
                    message[3],
                    message[4],
                    message[5],
                    message[6],
                    message[7],
                    message[8],
                    False,
                    False,
                    False,
                    chunk,
                    f"{row[6]}\n{chunk}",
                    embedding,
                    source_version,
                    digest,
                    json.dumps(metadata, sort_keys=True),
                    message[9],
                    True,
                    now,
                ],
            )
            source_revision = hashlib.sha256(
                json.dumps(
                    {
                        "file_id": str(row[4]),
                        "filename": row[6],
                        "mime_type": row[7],
                        "message_ts": row[9],
                        "size_bytes": path.stat().st_size,
                        "content_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                    },
                    sort_keys=True,
                ).encode()
            ).hexdigest()
            _set_projection_identity(
                connection,
                document_id,
                f"{source_version}:file:{row[4]}",
                source_revision,
                FILE_RECIPE_VERSION,
            )
            stats["chunks"] += 1
        stats["extracted"] += 1
    return stats


def project_documents(
    canonical_path: Path, state_path: Path, settings: Settings | None = None
) -> ProjectionStats:
    """Project messages, bounded thread contexts, and safe local file chunks."""
    if not canonical_path.exists():
        raise FileNotFoundError(canonical_path)
    gold_projection = settings is not None
    settings = settings or Settings(canonical_db=canonical_path, state_db=state_path)
    connection = connect_state(state_path)
    source_identity = str(canonical_path.resolve())
    try:
        connection.execute(f"ATTACH {path_sql(canonical_path)} AS canonical (READ_ONLY)")
        version_row = connection.execute(
            "SELECT max(schema_version) FROM canonical.schema_metadata"
        ).fetchone()
        assert version_row is not None
        if version_row[0] != CANONICAL_SCHEMA_VERSION:
            raise RuntimeError(f"unsupported canonical schema version: {version_row[0]}")
        watermark = _watermark(connection)
        now_row = connection.execute("SELECT current_timestamp").fetchone()
        assert now_row is not None
        now = now_row[0]
        connection.execute("BEGIN")
        connection.execute(
            """
            INSERT OR REPLACE INTO document_projection (
              document_id, document_kind, workspace_id, workspace_name, workspace_slug,
              workspace_url, channel_id, channel_name, author_id, author_name, ts, ts_us,
              timestamp, thread_id, thread_root_id, is_thread_parent, is_thread_reply,
              is_thread_broadcast, text_display, text_lexical, text_embedding,
              source_version, text_hash, metadata_json, permalink, is_active, projected_at
            )
            WITH canonical_users AS (
              SELECT workspace_id, user_id, any_value(username) AS username,
                     any_value(real_name) AS real_name, any_value(display_name) AS display_name
              FROM canonical.users GROUP BY workspace_id, user_id
            )
            SELECT m.message_key, 'message', m.workspace_id, w.team_name, w.workspace_slug,
              w.workspace_url, m.channel_id, c.name, m.user_id,
              coalesce(nullif(u.display_name, ''), nullif(u.real_name, ''), u.username),
              m.ts, m.ts_us, to_timestamp(m.ts_us / 1000000.0),
              CASE WHEN m.thread_ts IS NOT NULL OR m.is_thread_parent
                THEN m.workspace_id || ':' || m.channel_id || ':' ||
                     coalesce(m.thread_ts, m.ts) ELSE NULL END,
              CASE WHEN m.thread_ts IS NOT NULL OR m.is_thread_parent
                THEN coalesce(root.message_key, m.message_key) ELSE NULL END,
              m.is_thread_parent, m.is_thread_reply, m.is_thread_broadcast,
              coalesce(m.text, ''), m.search_text,
              concat('workspace: ', w.workspace_slug,
                '\nchannel: ', coalesce(c.name, m.channel_id),
                '\nauthor: ', coalesce(nullif(u.display_name, ''),
                  nullif(u.real_name, ''), u.username, m.user_id, 'unknown'),
                '\nmessage: ', coalesce(m.text, '')),
              ?, m.payload_hash,
              json_object('subtype', m.subtype, 'is_edited', m.is_edited),
              CASE WHEN w.workspace_url IS NULL THEN NULL ELSE
                rtrim(w.workspace_url, '/') || '/archives/' || m.channel_id || '/p' ||
                replace(m.ts, '.', '') || CASE WHEN m.is_thread_reply THEN
                  '?thread_ts=' || m.thread_ts || '&cid=' || m.channel_id ELSE '' END END,
              (NOT m.is_deleted AND
               nullif(trim(coalesce(m.search_text, m.text, '')), '') IS NOT NULL), ?
            FROM canonical.messages m
            JOIN canonical.workspaces w USING (workspace_id)
            LEFT JOIN canonical.channels c USING (workspace_id, channel_id)
            LEFT JOIN canonical_users u ON u.workspace_id=m.workspace_id AND u.user_id=m.user_id
            LEFT JOIN canonical.messages root ON root.workspace_id=m.workspace_id
              AND root.channel_id=m.channel_id AND root.ts=m.thread_ts
            """,
            [source_identity, now],
        )
        connection.execute(
            """
            UPDATE document_projection
            SET source_identity=?, source_revision=source_version, embedding_recipe=?
            WHERE document_kind='message'
              AND document_id IN (SELECT message_key FROM canonical.messages)
            """,
            [source_identity, MESSAGE_RECIPE_VERSION],
        )
        connection.execute(
            """
            UPDATE document_embeddings e SET source_version=p.source_version
            FROM document_projection p
            WHERE e.document_id=p.document_id AND e.text_hash=p.text_hash
              AND e.state='succeeded' AND e.embedding IS NOT NULL
              AND p.document_kind='message' AND p.embedding_recipe=?
            """,
            [MESSAGE_RECIPE_VERSION],
        )
        connection.execute(
            """
            UPDATE document_projection p SET is_active=false, projected_at=?
            WHERE p.document_kind='message' AND NOT EXISTS (
              SELECT 1 FROM canonical.messages m
              WHERE m.message_key=p.document_id AND NOT m.is_deleted
                AND nullif(trim(coalesce(m.search_text, m.text, '')), '') IS NOT NULL
            )
            """,
            [now],
        )
        connection.execute(
            "UPDATE document_projection SET is_active=false, projected_at=? "
            "WHERE source_identity=? AND document_kind='thread_context'",
            [now, source_identity],
        )
        thread_count = (
            _project_threads(connection, source_identity, now, settings) if gold_projection else 0
        )
        connection.execute(
            "UPDATE document_projection SET is_active=false, projected_at=? "
            "WHERE source_identity LIKE ? AND document_kind='file_chunk'",
            [now, f"{source_identity}:file:%"],
        )
        extraction = (
            _project_files(connection, source_identity, now, settings)
            if gold_projection
            else Counter()
        )
        connection.execute(
            """
            INSERT OR REPLACE INTO projection_watermarks VALUES (?, ?, ?, ?,
              (SELECT count(*) FROM document_projection WHERE source_identity=? AND is_active))
        """,
            [source_identity, CANONICAL_SCHEMA_VERSION, watermark, now, source_identity],
        )
        kind_rows = connection.execute(
            """
            SELECT document_kind, count(*), count(*) FILTER (WHERE is_active)
            FROM document_projection WHERE source_identity=? OR source_identity LIKE ?
            GROUP BY document_kind
        """,
            [source_identity, f"{source_identity}:file:%"],
        ).fetchall()
        projected_by_kind = {row[0]: row[1] for row in kind_rows}
        active_by_kind = {row[0]: row[2] for row in kind_rows}
        connection.execute(
            "INSERT OR REPLACE INTO projection_reports VALUES (?, ?, ?, ?, ?, ?)",
            [
                source_identity,
                watermark,
                now,
                json.dumps(projected_by_kind, sort_keys=True),
                json.dumps(active_by_kind, sort_keys=True),
                json.dumps({**dict(extraction), "thread_contexts": thread_count}, sort_keys=True),
            ],
        )
        connection.execute("COMMIT")
        projected = sum(projected_by_kind.values())
        active = sum(active_by_kind.values())
        return ProjectionStats(
            projected=projected,
            active=active,
            tombstoned=projected - active,
            watermark=watermark,
            projected_by_kind=projected_by_kind,
            active_by_kind=active_by_kind,
            extraction={**dict(extraction), "thread_contexts": thread_count},
        )
    except Exception:
        with suppress(Exception):
            connection.execute("ROLLBACK")
        raise
    finally:
        connection.close()
