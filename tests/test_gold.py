from __future__ import annotations

import fcntl
import os
import time
from pathlib import Path

import duckdb
import pytest
from docx import Document
from pptx import Presentation
from starlette.testclient import TestClient

from slackquery.artifact import (
    _lock_state,
    backup_state,
    restore_state,
    retain_artifacts,
    rollback_artifact,
)
from slackquery.db import connect_state
from slackquery.projection import (
    _safe_attachment_path,
    chunk_text,
    extract_file_text,
    project_documents,
)
from slackquery.retrieval import _exact_boost, route_query, weighted_rrf
from slackquery.server import create_asgi_app
from slackquery.settings import Settings


def test_extract_text_csv_docx_pptx_and_pdf(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    text = tmp_path / "sample.py"
    text.write_text("print('gold')\n")
    csv = tmp_path / "sample.csv"
    csv.write_text("name,value\r\ngold,1\r\n")
    docx = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("Gold document")
    document.save(docx)
    pptx = tmp_path / "sample.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Gold slides"
    presentation.save(pptx)
    pdf = tmp_path / "sample.pdf"
    pdf.write_bytes(b"%PDF-test")

    class Page:
        def extract_text(self) -> str:
            return "Gold PDF"

    class Reader:
        def __init__(self) -> None:
            self.pages = [Page()]

    monkeypatch.setattr("slackquery.projection.PdfReader", lambda _: Reader())
    assert extract_file_text(text, "text/x-python", 10_000) == "print('gold')"
    assert extract_file_text(csv, "text/csv", 10_000) == "name,value\ngold,1"
    assert extract_file_text(docx, None, 100_000) == "Gold document"
    assert extract_file_text(pptx, None, 100_000) == "Gold slides"
    assert extract_file_text(pdf, "application/pdf", 10_000) == "Gold PDF"


def test_attachment_safety_binary_oversize_and_chunks(tmp_path: Path) -> None:
    root = tmp_path / "attachments"
    workspace = root / "acme"
    workspace.mkdir(parents=True)
    safe = workspace / "safe.txt"
    safe.write_text("safe")
    outside = tmp_path / "outside.txt"
    outside.write_text("outside")
    (workspace / "escape.txt").symlink_to(outside)
    assert _safe_attachment_path(root, "acme", "safe.txt") == safe
    for value in ("../outside.txt", str(outside), "escape.txt"):
        with pytest.raises((ValueError, FileNotFoundError)):
            _safe_attachment_path(root, "acme", value)
    binary = workspace / "binary.bin"
    binary.write_bytes(b"x\x00y")
    with pytest.raises((TypeError, ValueError)):
        extract_file_text(binary, "text/plain", 100)
    with pytest.raises(OverflowError):
        extract_file_text(safe, "text/plain", 2)
    value = "abcdefghijklmnopqrstuvwxyz"
    chunks = chunk_text(value, target_chars=10, overlap_chars=2)
    assert chunks == chunk_text(value, target_chars=10, overlap_chars=2)
    assert chunks[0][-2:] == chunks[1][:2]


def test_gold_projection_thread_files_edit_and_delete(settings: Settings, canonical: Path) -> None:
    workspace = settings.attachment_root / "acme"
    workspace.mkdir(parents=True)
    (workspace / "notes.md").write_text("Gold attachment text " * 200)
    (workspace / "bad.bin").write_bytes(b"bad\x00binary")
    connection = duckdb.connect(str(canonical))
    connection.execute(
        """
        CREATE TABLE files(
          workspace_id VARCHAR, file_id VARCHAR, local_object_path VARCHAR,
          name VARCHAR, mimetype VARCHAR, message_key VARCHAR, channel_id VARCHAR,
          is_deleted BOOLEAN, payload_hash VARCHAR
        );
        INSERT INTO files VALUES
          ('W1', 'F1', 'notes.md', 'notes.md', 'text/markdown',
           'W1:C1:1.000001', 'C1', false, 'f1'),
          ('W1', 'F2', '../escape.txt', 'escape.txt', 'text/plain', NULL, 'C1', false, 'f2'),
          ('W1', 'F3', 'missing.txt', 'missing.txt', 'text/plain', NULL, 'C1', false, 'f3'),
          ('W1', 'F4', 'bad.bin', 'bad.bin', 'application/octet-stream', NULL, 'C1', false, 'f4');
        """
    )
    connection.close()
    configured = settings.model_copy(
        update={"file_chunk_tokens": 32, "file_chunk_overlap_tokens": 4}
    )
    first = project_documents(canonical, settings.state_db, configured)
    assert first.active_by_kind["message"] == 3
    assert first.active_by_kind["thread_context"] == 1
    assert first.active_by_kind["file_chunk"] > 1
    assert first.extraction["extracted"] == 1
    assert first.extraction["missing"] == 1
    assert first.extraction["rejected"] == 2
    state = connect_state(settings.state_db)
    old = state.execute(
        "SELECT text_hash FROM document_projection WHERE document_kind='thread_context'"
    ).fetchone()[0]
    state.close()
    canonical_db = duckdb.connect(str(canonical))
    canonical_db.execute(
        "UPDATE messages SET text='edited reply', search_text='edited reply', "
        "payload_hash='edited' WHERE message_key='W1:C1:2.000001'"
    )
    canonical_db.close()
    project_documents(canonical, settings.state_db, configured)
    state = connect_state(settings.state_db)
    edited = state.execute(
        "SELECT text_hash, text_display FROM document_projection "
        "WHERE document_kind='thread_context'"
    ).fetchone()
    assert edited[0] != old and "edited reply" in edited[1]
    state.close()


def test_projection_reuses_only_completed_message_embeddings(
    settings: Settings, canonical: Path
) -> None:
    project_documents(canonical, settings.state_db)
    state = connect_state(settings.state_db)
    rows = state.execute(
        "SELECT document_id, source_version, text_hash "
        "FROM document_projection ORDER BY document_id"
    ).fetchall()
    succeeded = rows[0]
    failed = rows[1]
    state.executemany(
        """
        INSERT INTO document_embeddings
          (document_id, generation_id, source_version, text_hash, embedding,
           native_dimension, state, attempt_count, updated_at)
        VALUES (?, ?, ?, ?, ?, 512, ?, 1, current_timestamp)
        """,
        [
            [succeeded[0], settings.generation_id, *succeeded[1:], [1.0] + [0.0] * 511,
             "succeeded"],
            [failed[0], settings.generation_id, *failed[1:], None, "retryable_failed"],
        ],
    )
    state.close()

    project_documents(canonical, settings.state_db, settings)
    state = connect_state(settings.state_db)
    try:
        source_identity = str(canonical.resolve())
        embedding_rows = state.execute(
            """
            SELECT document_id, source_version, state
            FROM document_embeddings
            WHERE document_id IN (?, ?)
            ORDER BY document_id
            """,
            [succeeded[0], failed[0]],
        ).fetchall()
        assert embedding_rows == [
            (succeeded[0], source_identity, "succeeded"),
            (failed[0], failed[1], "retryable_failed"),
        ]
    finally:
        state.close()


def test_actual_files_schema_workspace_name_and_message_ts(
    settings: Settings, canonical: Path
) -> None:
    workspace = settings.attachment_root / "Acme"
    path = workspace / "__uploads" / "F1" / "actual.txt"
    path.parent.mkdir(parents=True)
    path.write_text("actual canonical attachment")
    connection = duckdb.connect(str(canonical))
    connection.execute(
        """
        CREATE TABLE files(
          workspace_id VARCHAR, file_id VARCHAR, channel_id VARCHAR,
          message_ts VARCHAR, filename VARCHAR, mime_type VARCHAR,
          local_object_path VARCHAR
        );
        INSERT INTO files VALUES
          ('W1', 'F1', 'C1', '1.000001', 'actual.txt', 'text/plain',
           '__uploads/F1/actual.txt');
        """
    )
    connection.close()
    report = project_documents(canonical, settings.state_db, settings)
    assert report.extraction["extracted"] == 1
    state = connect_state(settings.state_db)
    row = state.execute(
        "SELECT channel_id, ts, metadata_json FROM document_projection "
        "WHERE document_kind='file_chunk' AND is_active"
    ).fetchone()
    state.close()
    assert row is not None
    assert row[:2] == ("C1", "1.000001")
    metadata = row[2] if isinstance(row[2], dict) else __import__("json").loads(row[2])
    assert metadata["message_id"] == "W1:C1:1.000001"


def test_routing_weighted_rrf_and_exact_boost() -> None:
    assert route_query("ERR-42") == "exact"
    assert route_query("how did the database incident unfold over time") == "conceptual"
    assert route_query("database ERR-42 investigation") == "mixed"
    rankings = {"lexical": ["a", "b"], "semantic": ["b", "a"]}
    lexical = weighted_rrf(rankings, {"lexical": 2.0, "semantic": 0.1}, 60)
    semantic = weighted_rrf(rankings, {"lexical": 0.1, "semantic": 2.0}, 60)
    assert lexical["a"] > lexical["b"]
    assert semantic["b"] > semantic["a"]
    assert lexical == weighted_rrf(rankings, {"lexical": 2.0, "semantic": 0.1}, 60)
    assert _exact_boost("ERR-42", "failure ERR-42 here", 1.0, 0.1) > _exact_boost(
        "ERR-42", "unrelated", 1.0, 0.1
    )


def test_metrics_retention_backup_and_restore(settings: Settings) -> None:
    with TestClient(create_asgi_app(settings)) as client:
        response = client.get("/metrics")
        assert response.status_code == 200
        assert "slackquery_requests_total" in response.text
        assert 'slackquery_backend_info{mode="ollama"} 1' in response.text
        assert "slackquery_route_requests_total" in response.text
        assert "slackquery_request_latency_seconds_bucket" in response.text
    settings.artifact_dir.mkdir(parents=True, exist_ok=True)
    artifacts = []
    for index in range(5):
        artifact = settings.artifact_dir / f"search-{index}.duckdb"
        artifact.touch()
        old = time.time() - 10_000 - index
        os.utime(artifact, (old, old))
        artifacts.append(artifact)
    retained = settings.model_copy(update={"retain_artifacts": 2, "retention_min_age_seconds": 0})
    report = retain_artifacts(retained)
    assert len(report["kept"]) == 2
    assert len(report["removed"]) == 3
    state = connect_state(settings.state_db)
    state.close()
    backup = settings.state_db.with_name("backup.duckdb")
    assert backup_state(settings, backup)["table_count"] > 0
    settings.state_db.unlink()
    assert restore_state(settings, backup)["valid"] is True
    restored = duckdb.connect(str(settings.state_db), read_only=True)
    assert (
        restored.execute(
            "SELECT count(*) FROM information_schema.tables WHERE table_name='document_projection'"
        ).fetchone()[0]
        == 1
    )
    restored.close()


def test_backup_restore_refuse_maintenance_lock(settings: Settings) -> None:
    state = connect_state(settings.state_db)
    state.close()
    lock = _lock_state(settings)
    try:
        with pytest.raises(RuntimeError, match="in use"):
            backup_state(settings, settings.state_db.with_name("busy-backup.duckdb"))
        with pytest.raises(RuntimeError, match="in use"):
            restore_state(settings, settings.state_db)
    finally:
        fcntl.flock(lock.fileno(), fcntl.LOCK_UN)
        lock.close()


def test_rollback_rejects_escape_and_invalid_artifact(settings: Settings, tmp_path: Path) -> None:
    outside = tmp_path / "outside.duckdb"
    outside.touch()
    with pytest.raises(ValueError, match="artifact_dir"):
        rollback_artifact(settings, str(outside))
    settings.artifact_dir.mkdir(parents=True, exist_ok=True)
    invalid = settings.artifact_dir / "search-invalid.duckdb"
    invalid.touch()
    with pytest.raises(duckdb.Error):
        rollback_artifact(settings, "invalid")


def test_restore_rejects_non_state_database(settings: Settings, tmp_path: Path) -> None:
    invalid = tmp_path / "invalid.duckdb"
    connection = duckdb.connect(str(invalid))
    connection.execute("CREATE TABLE unrelated(value INTEGER)")
    connection.close()
    with pytest.raises(RuntimeError, match="required state tables"):
        restore_state(settings, invalid)
